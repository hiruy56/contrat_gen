from fastapi import FastAPI, HTTPException, File, UploadFile
from docx import Document
from starlette.responses import FileResponse
import os

app = FastAPI()

def fill_contract_template(contract_details, output_path):
    template_path = os.path.abspath('tomp.docx')

    # Load the template document
    doc = Document(template_path)

    # Replace placeholders with actual values
    for paragraph in doc.paragraphs:
        for key, value in contract_details.items():
            if f"{{{key}}}" in paragraph.text:
                paragraph.text = paragraph.text.replace(f"{{{key}}}", str(value))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in contract_details.items():
                        if f"{{{key}}}" in paragraph.text:
                            paragraph.text = paragraph.text.replace(f"{{{key}}}", str(value))

    # Save the filled contract to the output path
    doc.save(output_path)

@app.post("/fill-contract")
async def fill_contract(contract_details: dict):
    output_path = os.path.abspath('filled_contract.docx')

    fill_contract_template(contract_details, output_path)

    return FileResponse(output_path, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
